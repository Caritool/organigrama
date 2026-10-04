"""Vuelca un .xlsx a JSON con la forma que Office Scripts entrega a Publicar.

Cada hoja sale como {"nombre", "valores"}, con `valores` igual a lo que daría
`getUsedRange(true).getValues()`: filas de celdas string/number/boolean, "" en
las vacías. Sirve para correr el núcleo de excel/publicar.ts fuera del Excel.

Uso: uv run --with openpyxl python tools/hojas-desde-xlsx.py libro.xlsx > hojas.json
"""
import json
import sys

from openpyxl import load_workbook


def celda(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor if isinstance(valor, (str, int, float, bool)) else str(valor)


libro = load_workbook(sys.argv[1], data_only=True)
hojas = []
for hoja in libro.worksheets:
    filas = [[celda(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]
    while filas and all(v == "" for v in filas[-1]):
        filas.pop()
    hojas.append({"nombre": hoja.title, "valores": filas})
json.dump(hojas, sys.stdout, ensure_ascii=False)
